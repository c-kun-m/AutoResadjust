from pathlib import Path
from docx import Document
from pypdf import PdfReader
import re,json,sys,hashlib

root=Path(__file__).resolve().parents[1];qa=root/'tmp/module_design_qa'
folder=qa/(sys.argv[1] if len(sys.argv)>1 else 'revised')
r=PdfReader(next(folder.glob('*.pdf')))
source=(root/'docs/工业级系统模块设计与协作方案.md').read_text('utf-8')
build=json.loads((qa/'build_report.json').read_text('utf-8'))
doc=Document(root/'output/跨数据中心异构AI算力平台模块设计与协作方案.docx')
norm=lambda s:re.sub(r'\s+','',s)
texts=[p.extract_text() for p in r.pages];full=norm('\n'.join(texts))
missing=[]
for p in doc.paragraphs:
    if len(p.text)>20 and '\t' not in p.text and norm(p.text) not in full:missing.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if norm(cell.text) not in full:missing.append(cell.text)
mp={}
def scan(items):
    for x in items:
        if isinstance(x,list):scan(x)
        else:mp[str(x.title)]=r.get_destination_page_number(x)+1
scan(r.outline)
pm={k:mp[k] for k in build['main_headings']}
previous=json.loads((qa/'page_map.json').read_text('utf-8')) if (qa/'page_map.json').exists() else {}
report={'pages':len(texts),'logical_pages':build['logical_pages'],'figures':len(doc.inline_shapes),'module_count':len(re.findall(r'^### M\d{2} ',source,re.M)), 'missing_text':missing,'tables':len(doc.tables),'toc_current':pm==previous,'page_map':pm,'page_summary':[{'page':i+1,'chars':len(t),'start':t[:90],'end':t[-75:]} for i,t in enumerate(texts)]}
assert not missing,missing
assert len(texts)==build['logical_pages'],report['page_summary']
assert report['module_count']==16 and report['figures']==5
assert not re.search(r'```|CREATE TABLE|/api/|数据库表|接口清单|TODO|TBD',source)
assert len(doc.element.xpath('.//w:pBdr'))==0
assert len(list(folder.glob('page-*.png')))==len(texts)
(qa/'page_map.json').write_text(json.dumps(pm,ensure_ascii=False,indent=2),'utf-8')
(qa/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
(qa/'extracted.txt').write_text('\n\n'.join('PAGE '+str(i+1)+'\n'+t for i,t in enumerate(texts)),'utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

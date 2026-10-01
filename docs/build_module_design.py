from pathlib import Path
import re, json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
import build_overview as base

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'docs/工业级系统模块设计与协作方案.md'
OUT=ROOT/'output/跨数据中心异构AI算力平台模块设计与协作方案.docx'
QA=ROOT/'tmp/module_design_qa'
QA.mkdir(parents=True,exist_ok=True)

class Diagram:
    def __init__(self,w=1800,h=900):
        self.im=Image.new('RGB',(w,h),'white');self.d=ImageDraw.Draw(self.im)
        self.ft=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',30)
        self.fb=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc',34)
        self.fs=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',27)
    def box(self,xy,title,lines=(),fill='#F0F5F9'):
        x0,y0,x1,y1=xy;self.d.rounded_rectangle(xy,12,fill,outline='#8E9FAA',width=2)
        total=41+len(lines)*39;cy=(y0+y1-total)/2
        self.d.text(((x0+x1)/2,cy),title,font=self.fb,anchor='mt',fill='black')
        for i,l in enumerate(lines):self.d.text(((x0+x1)/2,cy+46+i*39),l,font=self.ft,anchor='mt',fill='#222222')
    def arrow(self,pts,color='#526C83',width=4):
        self.d.line(pts,fill=color,width=width)
        x,y=pts[-1];px,py=pts[-2];a=math.atan2(y-py,x-px)
        self.d.polygon([(x,y),(x-20*math.cos(a-.45),y-20*math.sin(a-.45)),(x-20*math.cos(a+.45),y-20*math.sin(a+.45))],fill=color)
    def label(self,xy,txt,anchor='mm',color='#333333'):
        self.d.text(xy,txt,font=self.fs,fill=color,anchor=anchor)
    def save(self,name):self.im.save(ROOT/'docs'/name)

def figures():
    d=Diagram(1800,1110)
    d.box((25,20,1400,125),'M01 统一门户与工作空间')
    d.box((25,190,680,300),'M02 租户身份与权限')
    d.box((735,190,1400,300),'M03 配额与服务等级')
    d.arrow([(370,125),(370,190)]);d.arrow([(1050,125),(1050,190)])
    for xy,t,l in [((25,365,435,535),'资源与环境',['M04 资源中心','M05 异构运行环境']),((480,365,945,535),'作业与放置',['M06 任务编排','M07 全局调度']),((990,365,1400,535),'M11 推理服务管理',['版本  副本  容量'])]:d.box(xy,t,l)
    d.arrow([(360,300),(360,335),(710,335),(710,365)])
    d.arrow([(1070,300),(1070,335),(1195,335),(1195,365)])
    d.box((25,605,945,775),'模型与数据协同',['M08 模型与数据资产','M09 传输与缓存'])
    d.box((990,605,1400,775),'M12 在线请求路由',['独立调用入口','授权  流控  路由'])
    d.arrow([(230,535),(230,570),(485,570),(485,605)])
    d.arrow([(715,535),(715,605)])
    d.arrow([(1195,535),(1195,605)])
    d.box((25,855,1400,1030),'M10 机房执行与自治',['机房 A     机房 B     机房 N','本地准入与拓扑    训练与批推    已就绪推理副本'])
    d.arrow([(480,775),(480,855)],'#AD7638',6)
    d.arrow([(1195,775),(1195,855)],'#147978',7)
    d.arrow([(945,450),(968,450),(968,855)])
    d.arrow([(1400,450),(1430,450),(1430,945),(1400,945)])
    d.box((1470,20,1780,1030),'',[],fill='#F5F6F7')
    d.d.text((1625,65),'生产运营',font=d.fb,anchor='mt',fill='black')
    for y,t,ls in [(200,'M13',['观测与运维']),(390,'M14',['计量与成本']),(580,'M15',['安全与治理']),(770,'M16',['交付与连续性'])]:
        d.d.text((1625,y),t,font=d.fb,anchor='mt',fill='black');d.label((1625,y+65),ls[0])
    d.label((30,1070),'灰蓝线  管理控制       橙色线  资产传输       绿色线  在线请求',anchor='lm')
    d.save('module_architecture.png')

    d=Diagram(1800,790)
    positions=[(20,30,515,180),(650,30,1145,180),(1280,30,1775,180),(1280,310,1775,475),(650,310,1145,475),(20,310,515,475)]
    names=[('1 申请与校验',['M01 → M06','授权  权益  环境  版本']),('2 候选选址',['M07','资源  拓扑  数据位置']),('3 数据准备',['M09','版本送达并通过校验']),('4 容量最终确认',['M03 与 M10','配额  设备组  本地条件']),('5 成组执行',['M10 → M06','运行状态  检查点']),('6 产物与释放确认',['M08 与 M10','结果可用  资源已释放'])]
    for xy,(t,l) in zip(positions,names):d.box(xy,t,l)
    d.arrow([(515,105),(650,105)]);d.arrow([(1145,105),(1280,105)]);d.arrow([(1525,180),(1525,310)])
    d.arrow([(1280,395),(1145,395)]);d.arrow([(650,395),(515,395)])
    d.box((20,595,1775,735),'贯穿执行的生产支撑',['M13 观测与诊断    M14 用量与对账    M15 授权和安全约束'])
    d.label((900,535),'准备失败重做准备    准入失败重新排队    执行恢复先确认旧执行')
    d.save('training_flow.png')

    d=Diagram(1800,865)
    d.label((30,35),'服务部署链路',anchor='lm')
    xs=[25,475,925,1375]
    ns=[('M11 服务管理',['版本与容量要求']),('M07 全局调度',['选址与资源计划']),('M09 与 M10',['模型准备与部署']),('M11 就绪确认',['预热与业务探测'])]
    for x,(t,l) in zip(xs,ns):d.box((x,90,x+395,245),t,l)
    for x in xs[:-1]:d.arrow([(x+395,168),(x+450,168)])
    d.arrow([(1570,245),(1570,315),(690,315),(690,410)])
    d.label((1160,282),'发布有效副本与流量规则')
    d.label((30,365),'在线请求链路',anchor='lm')
    d.box((25,410,380,580),'业务调用方',['请求与流式响应'])
    d.box((485,410,905,580),'M12 在线请求路由',['授权  限流  选择副本'])
    d.box((1010,410,1420,580),'机房内已就绪副本',['加载完毕的模型实例'])
    d.arrow([(380,480),(485,480)],'#147978',6);d.arrow([(905,480),(1010,480)],'#147978',6)
    d.arrow([(1010,540),(905,540)],'#147978',6);d.arrow([(485,540),(380,540)],'#147978',6)
    d.box((25,680,1770,815),'容量与运营反馈',['M13 汇总排队和延迟 → M11 组织扩容    M14 异步汇集可信调用用量'])
    d.arrow([(1210,580),(1210,680)])
    d.save('inference_flow.png')

    d=Diagram(1800,840)
    d.box((530,15,1270,135),'中心无法取得机房新状态',['先标记不确定并停止新增投放'])
    d.box((25,245,835,435),'机房本地继续运行',['M10 维持已授权执行与受限恢复','状态和用量缓冲    既有占用保留'])
    d.box((965,245,1775,435),'中心管理迁移风险',['M06 核查旧执行    M12 独立探测服务','唯一执行作业不可因失联直接重发'])
    d.arrow([(760,135),(760,190),(430,190),(430,245)])
    d.arrow([(1040,135),(1040,190),(1370,190),(1370,245)])
    d.box((25,570,1775,775),'重连后按顺序核对',['确认控制权 → 核对运行与设备 → 修复状态差异 → 补传用量 → 逐步恢复投放','需要迁移时  先确认旧执行结束或完成有效隔离  再启动新的执行'])
    d.arrow([(430,435),(430,570)]);d.arrow([(1370,435),(1370,570)])
    d.save('recovery_flow.png')

    d=Diagram(1800,1040)
    d.box((25,20,1275,215),'全局中心集群',['业务治理    任务编排与调度    资源与资产目录','观测    用量    安全    配置与交付'])
    d.box((1350,20,1775,215),'灾备中心',['恢复材料与预案','受控接管'])
    d.arrow([(1275,118),(1350,118)])
    d.box((25,315,850,500),'机房 A 执行单元',['代理与本地控制    缓存    观测与用量缓冲','已认证资源池    训练与批推    推理副本'])
    d.box((950,315,1775,500),'机房 B 执行单元',['代理与本地控制    缓存    观测与用量缓冲','已认证资源池    训练与批推    推理副本'])
    d.arrow([(350,215),(350,315)]);d.arrow([(1060,215),(1060,260),(1370,260),(1370,315)])
    d.box((25,615,1775,760),'按地域分组的在线入口',['独立承载请求    使用已发布规则    仅选择符合约束的就绪副本'])
    d.arrow([(435,615),(435,500)],'#147978',6);d.arrow([(1370,615),(1370,500)],'#147978',6)
    d.box((25,855,1775,1000),'隔离与扩展原则',['中心按职责扩容    机房按资源池隔离    在线入口按服务与区域扩容','开发验证环境独立    关键状态受保护    恢复先隔离旧控制权'])
    d.save('deployment_architecture.png')

def rich(p,txt):
    for chunk in re.split(r'(\*\*.*?\*\*|\[R\d{2}\])',txt):
        if chunk.startswith('**') and chunk.endswith('**'):base.font(p.add_run(chunk[2:-2]),bold=True)
        elif re.fullmatch(r'\[R\d{2}\]',chunk):base.add_link(p,chunk,anchor=chunk[1:-1])
        else:base.font(p.add_run(chunk))

def para(doc,txt):
    p=doc.add_paragraph();rich(p,txt);return p

def make_table(doc,rows):
    t=base.table(doc,rows)
    widths=[1.25,2.45,3.05]
    if rows[0][0]=='分组':widths=[.85,2.20,3.70]
    if rows[0][0]=='联合场景':widths=[1.5,3.25,2.0]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j])
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.13
                p.paragraph_format.keep_with_next=(i<2)
                for r in p.runs:base.font(r,9.5,bold=i==0)
    return t

def main():
    figures();doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.left_margin=sec.right_margin=Inches(.76);sec.top_margin=sec.bottom_margin=Inches(.70)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.30)
    for n in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[n];s.font.name=base.FONT;s.font.color.rgb=RGBColor(0,0,0);s.font.italic=False
        rf=s._element.get_or_add_rPr().rFonts
        for a in list(rf.attrib):
            if 'theme' in a.lower():del rf.attrib[a]
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),base.FONT)
    for e in doc.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.38;normal.paragraph_format.space_after=Pt(9)
    normal.paragraph_format.widow_control=True
    for n,size,be,af in [('Title',25,0,18),('Subtitle',17,0,18),('Heading 1',18,0,14),('Heading 2',14,10,10)]:
        s=doc.styles[n];s.font.size=Pt(size);s.paragraph_format.space_before=Pt(be);s.paragraph_format.space_after=Pt(af)
        s.paragraph_format.keep_with_next=True
    sec.different_first_page_header_footer=True
    base.font(sec.header.paragraphs[0].add_run('跨数据中心异构 AI 算力管理与调度平台'),8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    base.font(footer.add_run('模块设计与协作方案    2.0     '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    p=doc.add_paragraph(style='Title');base.font(p.add_run('跨数据中心异构 AI 算力\n管理与调度平台'),25,bold=True)
    doc.add_paragraph('工业级系统模块设计与协作方案',style='Subtitle')
    para(doc,'版本 2.0    2026 年 9 月 19 日')
    para(doc,'适用对象    项目负责人  产品与架构团队  研发  运维  安全与运营')
    doc.add_heading('系统设计结论',1)
    para(doc,'采用全局统筹与机房自治的分层结构，围绕资源、任务、模型和在线服务建立明确的模块分工，以完整业务流程和故障恢复组织模块协作。')
    para(doc,'系统由 16 个一级模块组成，覆盖业务治理、资源调度、模型与执行、在线推理和生产运营。每个模块定义内部能力、责任边界、协作对象和生产运行要求。')
    doc.add_heading('阅读重点',1)
    for txt in ['第 1 至 3 章    系统目标  总体架构  模块清单','第 4 至 8 章    16 个模块的设计和相互关系','第 9 章    训练  推理  失联恢复  扩容维护与资产退役','第 10 至 11 章    生产部署  故障隔离  责任划分与联合验证']:para(doc,txt)
    para(doc,'工业级能力贯穿日常运行、故障处置、升级扩容和灾备接管，以长期稳定运营作为设计完成标准。')
    doc.add_page_break();doc.add_heading('目录',1)
    text=SRC.read_text('utf-8');headings=re.findall(r'^## (.+)$',text,re.M)
    pmfile=QA/'page_map.json';pm=json.loads(pmfile.read_text('utf-8')) if pmfile.exists() else {}
    for i,title in enumerate(headings):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10);p.paragraph_format.line_spacing=1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.72),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        base.add_link(p,title,anchor='s'+str(i));base.font(p.add_run('\t'+str(pm.get(title,''))),11)
        if title.startswith('4 '):para(doc,'M01 统一门户与工作空间    M02 租户身份与权限    M03 配额与服务等级')
        if title.startswith('5 '):para(doc,'M04 资源中心    M05 异构运行环境    M06 任务编排    M07 全局调度')
        if title.startswith('6 '):para(doc,'M08 模型与数据资产    M09 传输与缓存    M10 机房执行与自治')
        if title.startswith('7 '):para(doc,'M11 推理服务管理    M12 在线请求路由')
        if title.startswith('8 '):para(doc,'M13 可观测性与运维    M14 计量与成本运营\nM15 安全与治理    M16 平台交付与连续性')
    doc.add_page_break()
    lines=text.splitlines();i=next(i for i,l in enumerate(lines) if l.startswith('## '));idx=0;imcount=0
    logical_pages=3;new_page=False
    while i<len(lines):
        l=lines[i].strip()
        if not l:i+=1;continue
        if l=='<!-- page -->':new_page=True;logical_pages+=1
        elif l.startswith('## '):
            p=doc.add_heading(l[3:],1);base.bookmark(p,'s'+str(idx),idx+1);idx+=1
            if new_page:p.paragraph_format.page_break_before=True;new_page=False
        elif l.startswith('### '):
            p=doc.add_heading(l[4:],2)
            if new_page:p.paragraph_format.page_break_before=True;new_page=False
        elif l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                r=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in r):rows.append(r)
                i+=1
            make_table(doc,rows);continue
        elif l.startswith('!['):
            cap,path=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',l).groups();imcount+=1
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ROOT/'docs'/path),width=Inches(6.72))
            p._p.xpath('.//wp:docPr')[0].set('descr',cap)
            c=doc.add_paragraph();c.alignment=WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after=Pt(12);base.font(c.add_run(cap),9)
        elif l.startswith('https://'):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(15);base.add_link(p,l,url=l)
        else:
            p=para(doc,l)
            if re.match(r'^\[R\d{2}\]',l):
                rid=l[1:4];base.bookmark(p,rid,100+int(rid[1:]));p.paragraph_format.keep_with_next=True
        i+=1
    for p in doc.paragraphs:
        if p.style.name in ['Title','Subtitle','Heading 1','Heading 2']:
            size={'Title':25,'Subtitle':17,'Heading 1':18,'Heading 2':14}[p.style.name]
            for r in p.runs:base.font(r,size,bold=p.style.name!='Subtitle');r.font.italic=False
    for e in doc.element.xpath('.//w:pBdr'):e.getparent().remove(e)
    doc.core_properties.title='跨数据中心异构AI算力平台模块设计与协作方案'
    doc.core_properties.subject='工业级系统模块设计与协作方案'
    doc.core_properties.author='项目规划'
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    (QA/'build_report.json').write_text(json.dumps({'logical_pages':logical_pages,'characters':len(text),'main_headings':headings,'figures':imcount},ensure_ascii=False,indent=2),'utf-8')
    print(OUT);print('logical_pages',logical_pages,'chars',len(text),'figures',imcount)

if __name__=='__main__':main()

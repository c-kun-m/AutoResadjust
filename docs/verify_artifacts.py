from pathlib import Path
import json,re,hashlib
from docx import Document
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1]
src=(ROOT/'docs/跨数据中心异构AI算力平台项目总览.md').read_text('utf-8')
docx=ROOT/'output/跨数据中心异构AI算力平台项目总览.docx'
pdf=ROOT/'tmp/doc_qa/final/跨数据中心异构AI算力平台项目总览.pdf'
d=Document(docx);r=PdfReader(pdf)
norm=lambda x:re.sub(r'\s+','',x)
whole=norm('\n'.join(p.extract_text() for p in r.pages))
missing=[]
for p in d.paragraphs:
    if len(p.text)>25 and p.style.name not in ['Title'] and p.text not in [''] and not p.text.startswith('https://'):
        if norm(p.text) not in whole and '\t' not in p.text:missing.append(p.text[:100])
maps={}
def outline(items):
    for x in items:
        if isinstance(x,list):outline(x)
        else:maps[str(x.title)]=r.get_destination_page_number(x)+1
outline(r.outline)
saved=json.loads((ROOT/'tmp/doc_qa/page_map.json').read_text('utf-8'))
assert all(maps[k]==v for k,v in saved.items()),'TOC page mismatch'
assert len(set(re.findall(r'F\d{2}',src)))==11
assert len(set(re.findall(r'AT\d{2}',src)))==22
assert len(set(re.findall(r'WP\d{2}',src)))==9
assert len(d.element.xpath('.//w:pBdr'))==0
assert len(d.tables)==25
assert len(list((ROOT/'tmp/doc_qa/final').glob('page-*.png')))==len(r.pages)
assert not re.search(r'TODO|TBD|MVP|待补充',src)
report={'pages':len(r.pages),'tables':len(d.tables),'chapters':23,'appendices':2,'functional_modules':11,'work_packages':9,'acceptance_cases':22,'toc_verified':True,'paragraphs_missing_in_pdf':missing,'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest()}
(ROOT/'tmp/doc_qa/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
